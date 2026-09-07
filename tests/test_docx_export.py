import io
from docx import Document

import pytest
from httpx import AsyncClient, ASGITransport

from app.database import Database
from app.docx_generator import generate_resume_docx, generate_cover_letter_docx


# --- Unit tests for docx_generator ---

def test_generate_resume_docx():
    resume = "Jane Doe\njane@example.com\n\nEXPERIENCE\nSenior Engineer | Acme Corp | 2020-2024\n- Built cool stuff\n- Led team of 5"
    result = generate_resume_docx(resume)
    assert isinstance(result, bytes)
    assert len(result) > 0
    # DOCX files start with PK (zip format)
    assert result[:2] == b"PK"
    text = "\n".join(p.text for p in Document(io.BytesIO(result)).paragraphs)
    assert "Jane Doe" in text and "Acme Corp" in text and "Led team of 5" in text


def test_generate_resume_docx_with_name():
    resume = "Jane Doe\nSome content"
    result = generate_resume_docx(resume, name="Jane Doe")
    assert isinstance(result, bytes)
    assert len(result) > 0


def test_generate_cover_letter_docx():
    letter = "Dear Hiring Manager,\n\nI am writing to apply for the position.\n\nSincerely,\nJane Doe"
    result = generate_cover_letter_docx(letter, company="Acme", position="Engineer")
    assert isinstance(result, bytes)
    assert len(result) > 0
    assert result[:2] == b"PK"
    text = "\n".join(p.text for p in Document(io.BytesIO(result)).paragraphs)
    assert "Dear Hiring Manager" in text and "Jane Doe" in text


def test_generate_cover_letter_docx_no_metadata():
    letter = "Dear Sir,\n\nHello.\n\nRegards,\nJohn"
    result = generate_cover_letter_docx(letter)
    assert isinstance(result, bytes)
    assert len(result) > 0


def test_generate_resume_docx_empty():
    result = generate_resume_docx("")
    assert isinstance(result, bytes)
    assert len(result) > 0


# --- API integration tests ---

@pytest.fixture
async def app(tmp_path):
    from app.main import create_app
    application = create_app(db_path=str(tmp_path / "test.db"), testing=True)
    db = Database(str(tmp_path / "test.db"))
    await db.init()
    application.state.db = db
    yield application
    await db.close()


@pytest.fixture
async def client(app):
    transport = ASGITransport(app=app)
    async with AsyncClient(transport=transport, base_url="http://test") as ac:
        yield ac


async def _create_job_with_application(db, resume_text="Test resume", cover_letter="Test letter"):
    job_id = await db.insert_job(
        title="Engineer", company="TestCo", location="Remote",
        salary_min=None, salary_max=None, description="A job",
        url="https://example.com/job/1", posted_date=None,
        application_method="url", contact_email=None,
    )
    await db.insert_application(job_id)
    app_row = await db.get_application(job_id)
    await db.update_application(app_row["id"], tailored_resume=resume_text, cover_letter=cover_letter)
    return job_id


@pytest.mark.asyncio
async def test_download_resume_docx(client, app):
    job_id = await _create_job_with_application(app.state.db, resume_text="Jane Doe\nExperience\n- Did things")
    resp = await client.get(f"/api/jobs/{job_id}/resume.docx")
    assert resp.status_code == 200
    assert "application/vnd.openxmlformats-officedocument.wordprocessingml.document" in resp.headers["content-type"]
    assert resp.content[:2] == b"PK"


@pytest.mark.asyncio
async def test_download_cover_letter_docx(client, app):
    job_id = await _create_job_with_application(app.state.db, cover_letter="Dear Hiring Manager,\n\nHello.")
    resp = await client.get(f"/api/jobs/{job_id}/cover-letter.docx")
    assert resp.status_code == 200
    assert "application/vnd.openxmlformats-officedocument.wordprocessingml.document" in resp.headers["content-type"]


@pytest.mark.asyncio
async def test_download_resume_docx_no_job(client):
    resp = await client.get("/api/jobs/999/resume.docx")
    assert resp.status_code == 404


@pytest.mark.asyncio
async def test_download_cover_letter_docx_no_application(client, app):
    job_id = await app.state.db.insert_job(
        title="Engineer", company="TestCo", location="Remote",
        salary_min=None, salary_max=None, description="A job",
        url="https://example.com/job/2", posted_date=None,
        application_method="url", contact_email=None,
    )
    resp = await client.get(f"/api/jobs/{job_id}/cover-letter.docx")
    assert resp.status_code == 404
